import re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
from sel_data import *
import data30 as M
import budget as BUD

OUT = "../output/MediBec_Scenario_Selection.xlsx"
wb = Workbook()
F = "Arial"
BLUE = Font(name=F, color="0000FF", size=10); BLK = Font(name=F, size=10); GRN = Font(name=F, color="008000", size=10)
HDR = Font(name=F, bold=True, color="FFFFFF", size=10); TTL = Font(name=F, bold=True, size=14); BD = Font(name=F, bold=True, size=10)
SUB = Font(name=F, bold=True, size=11, color="1F3864")
HFILL = PatternFill("solid", fgColor="1F3864"); YEL = PatternFill("solid", fgColor="FFFF00")
thin = Side(style="thin", color="BFBFBF"); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="top")
GRF = PatternFill("solid", fgColor="D9EAD3"); AMB = PatternFill("solid", fgColor="FFF2CC"); RED = PatternFill("solid", fgColor="F4CCCC"); GRY = PatternFill("solid", fgColor="EDEDED")

def header(ws, row, cols, widths=None, col0=1):
    for j, c in enumerate(cols):
        x = ws.cell(row=row, column=col0 + j, value=c); x.font, x.fill, x.border = HDR, HFILL, BOX
        x.alignment = Alignment(wrap_text=True, vertical="center")
    if widths:
        for j, w in enumerate(widths):
            ws.column_dimensions[L(col0 + j)].width = w
def put(ws, r, c, v, font=BLK, fmt=None, wrap=True, fill=None):
    x = ws.cell(row=r, column=c, value=v); x.font, x.border = font, BOX
    if fmt: x.number_format = fmt
    if wrap: x.alignment = WRAP
    if fill: x.fill = fill
    return x
def title(ws, t, sub=None):
    ws["A1"] = t; ws["A1"].font = TTL
    if sub: ws["A2"] = sub; ws["A2"].font = BLK
    ws.sheet_view.showGridLines = False
def table(ws, r0, cols, widths, rows, fonts=None):
    header(ws, r0, cols, widths)
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            put(ws, r0 + 1 + i, j + 1, v, (fonts[j] if fonts else BLK))
    return r0 + 1 + len(rows)

# ================================================================ README
ws = wb.active; ws.title = "README"
title(ws, "MediBec - Cybersecurity Risk Scenario Selection (10-scenario portfolio)")
lines = [
    "Prepared 25 September 2026 under the Scenario Selection Instructions (20 steps, output Sections A-E). Assessment horizon: 12 months (Oct 2026 - Sep 2027).",
    "Source: MediBec Business Case v2.0b (fictional teaching case) and the earlier 30-scenario CyberRiskGuardian analysis. All scores and estimates are ANALYST JUDGEMENT - VALIDATION REQUIRED.",
    "Final scenario selection remains the responsibility of the human analyst and MediBec management (Critical rule 14).",
    "",
    "Sheets (output sections):",
    "  A_Context, A_AssetMap, A_Dependencies  - Section A: scope, objectives, asset map, single points of failure, materiality criteria, gaps.",
    "  B_Candidates                           - Section B: 45 candidate scenarios written as causal chains.",
    "  C_Screening                            - Section C: 12 criteria scored 1-5 (blue = editable), escalation flags, evidence, confidence, disposition.",
    "  C_Weights, C_Sensitivity               - Section C / step 20: four weighting schemes; family ranking recalculates live.",
    "  D_Portfolio, D_WhyTest, D_QuantPrep, D_Validation - Section D: final 10 scenarios (P01-P10), 'why this scenario' test, inputs for quantification, stakeholder validation.",
    "  E_Coverage, E_Quality                  - Section E: coverage matrix, portfolio quality review, blind spots, evidence plan.",
    "  Changes                                - how P01-P10 relate to S1-S30 of the earlier assessment.",
    "",
    "Scale: 1 = very low, 3 = moderate, 5 = very high. For Evidence strength, 5 = direct organizational evidence, 1 = assumption only.",
    "The total score is a screening aid, not a ranking rule (Instruction s.9). Dispositions: Select / Consolidate (into a P scenario) / Consider (explicit review) / Defer.",
    "Colours: blue text = input; black = formula or analyst text; green = link to another sheet.",
]
for i, t in enumerate(lines, 3):
    ws.cell(row=i, column=1, value=t).font = BLK
ws.column_dimensions["A"].width = 160

# ================================================================ A_Context
ws = wb.create_sheet("A_Context"); title(ws, "Section A - Assessment context")
r = table(ws, 3, ["Element", "Summary", "Basis"], [26, 110, 34], SCOPE)
ws.cell(row=r + 1, column=1, value="Critical objectives").font = SUB
r = table(ws, r + 2, ["ID", "Objective", "Critical processes"], None, OBJECTIVES)
ws.cell(row=r + 1, column=1, value="Materiality criteria used for selection (risk appetite and tolerance, Instruction s.4)").font = SUB
r = table(ws, r + 2, ["Dimension", "Scenario is material if it could..."], None, MATERIALITY)
ws.cell(row=r + 1, column=1, value="Assumptions and information gaps (Instruction s.1, s.15)").font = SUB
r = table(ws, r + 2, ["Gap / assumption", "How to validate"], None, GAPS)

ws = wb.create_sheet("A_AssetMap"); title(ws, "Section A - Business objective -> process -> primary asset -> supporting assets -> dependencies -> consequences")
table(ws, 3, ["Objective", "Business process", "Primary asset", "Supporting assets", "Dependencies", "Potential consequences"], [10, 34, 32, 42, 38, 36], ASSET_MAP)

ws = wb.create_sheet("A_Dependencies"); title(ws, "Section A - Dependencies, single points of failure and concentration (Instruction s.3)",
                                              "Peripheral = assets that receive less security oversight but reach critical systems.")
table(ws, 3, ["Dependency", "Type", "What depends on it", "Single point of failure / concentration", "Peripheral asset?", "Related candidates"], [36, 16, 40, 48, 22, 26], DEPENDENCIES)

# ================================================================ B_Candidates
ws = wb.create_sheet("B_Candidates"); title(ws, "Section B - Candidate scenario population (45 candidates)",
                                            "Generation strategy: TD top-down, BU bottom-up, TH threat-driven, GL generic localized, DP dependency-driven. Prior ID links to the 30-scenario assessment.")
rows = cand_rows()
header(ws, 3, ["ID", "Scenario name", "Causal scenario statement", "Primary asset / process", "Threat source / event", "Principal vulnerability / condition", "Primary consequence", "CIA", "Strategy", "Source ref", "Prior ID"],
       [6, 32, 80, 18, 20, 30, 24, 8, 9, 14, 8])
for i, c in enumerate(rows):
    rr = 4 + i
    for j, k in enumerate(["id", "name", "statement", "asset", "threat", "vuln", "consequence", "cia", "strategy", "ref", "prev"]):
        put(ws, rr, j + 1, c[k])
ws.freeze_panes = "C4"

# ================================================================ C_Weights
ww = wb.create_sheet("C_Weights"); title(ww, "Section C - Criteria weighting schemes (sensitivity, Instruction s.20)", "Blue cells are editable; screening and sensitivity recalculate.")
header(ww, 3, ["Code", "Criterion", "Question"] + list(WEIGHTS.keys()), [6, 28, 70, 10, 11, 12, 12])
QUEST = ["Does the scenario threaten an important objective or critical process?", "Does it affect high-value information, technology, services or infrastructure?",
         "Is the threat credible for this organization and sector?", "Are there plausible weaknesses or predisposing conditions?", "Could consequences materially affect the organization?",
         "Could it create important compliance, privacy, contractual or legal consequences?", "Could it disrupt critical services or create cascading failures?",
         "Does it involve important third parties or single points of failure?", "Is there organizational, technical, historical, sector or threat-intelligence evidence?",
         "Could material consequences develop rapidly?", "Is it uncertain whether existing controls would prevent or contain it?", "Would analysing it potentially change a management decision?"]
for i in range(12):
    put(ww, 4 + i, 1, CRIT_SHORT[i]); put(ww, 4 + i, 2, CRIT[i]); put(ww, 4 + i, 3, QUEST[i])
    for j, (k, w) in enumerate(WEIGHTS.items()):
        put(ww, 4 + i, 4 + j, w[i], BLUE, "0.0")
WCOL = {k: L(4 + j) for j, k in enumerate(WEIGHTS)}

# ================================================================ C_Screening
wc = wb.create_sheet("C_Screening"); title(wc, "Section C - Screening matrix (scores 1-5; blue = editable)",
                                           "Total and weighted scores are screening aids only. Escalation flags: " + "; ".join(f"{k} = {v}" for k, v in ESCALATION_CODES.items()))
cols = ["ID", "Scenario"] + CRIT_SHORT + ["Total (/60)"] + [f"Weighted - {k}" for k in WEIGHTS] + ["Rank (equal)", "Escalation flags", "# flags", "Evidence basis", "Selection confidence", "Disposition", "Family / target", "Rationale for scores and disposition"]
widths = [6, 34] + [5] * 12 + [8] + [9] * 4 + [7, 18, 6, 26, 10, 12, 22, 70]
header(wc, 3, cols, widths)
N = len(rows); R0 = 4; RN = R0 + N - 1
for i, c in enumerate(rows):
    r = R0 + i; cid = c["id"]
    put(wc, r, 1, cid); put(wc, r, 2, c["name"])
    for j, v in enumerate(SCORES[cid]):
        put(wc, r, 3 + j, v, BLUE, "0")
    put(wc, r, 15, f"=SUM(C{r}:N{r})", fmt="0")
    for j, k in enumerate(WEIGHTS):
        col = WCOL[k]
        put(wc, r, 16 + j, "=" + "+".join(f"{L(3+q)}{r}*C_Weights!${col}${4+q}" for q in range(12)), fmt="0.0")
    put(wc, r, 20, f"=RANK(O{r},$O${R0}:$O${RN})", fmt="0")
    esc = ESC.get(cid, "")
    put(wc, r, 21, esc or "-")
    put(wc, r, 22, f'=IF(U{r}="-",0,LEN(U{r})-LEN(SUBSTITUTE(U{r},",",""))+1)', fmt="0")
    ev = EVIDENCE_TYPE.get(cid) or {5: "Verified org evidence", 4: "Case documentation", 3: "External/sector + case", 2: "Assumption + sector", 1: "Assumption / unknown"}[SCORES[cid][8]]
    put(wc, r, 23, ev); put(wc, r, 24, CONF[cid], BLUE)
    d, t = DISP[cid]
    put(wc, r, 25, d, BLUE); put(wc, r, 26, t, BLUE); put(wc, r, 27, RATIONALE[cid])
wc.freeze_panes = "C4"
wc.conditional_formatting.add(f"C{R0}:N{RN}", ColorScaleRule(start_type="num", start_value=1, start_color="F8F8F8", mid_type="num", mid_value=3, mid_color="FFE699", end_type="num", end_value=5, end_color="F4B183"))
for val, fill in (("Select", GRF), ("Consolidate", AMB), ("Consider", PatternFill("solid", fgColor="DDEBF7")), ("Defer", GRY)):
    wc.conditional_formatting.add(f"Y{R0}:Y{RN}", CellIsRule(operator="equal", formula=[f'"{val}"'], fill=fill))
dv = DataValidation(type="list", formula1='"Select,Consolidate,Consider,Defer"', allow_blank=False); wc.add_data_validation(dv); dv.add(f"Y{R0}:Y{RN}")
dv2 = DataValidation(type="whole", operator="between", formula1="1", formula2="5"); wc.add_data_validation(dv2); dv2.add(f"C{R0}:N{RN}")
r = RN + 2
put(wc, r, 2, "Disposition counts", BD, wrap=False)
for j, d in enumerate(("Select", "Consolidate", "Consider", "Defer")):
    put(wc, r + 1 + j, 2, d); put(wc, r + 1 + j, 3, f'=COUNTIF($Y${R0}:$Y${RN},"{d}")', fmt="0")
put(wc, r + 5, 2, "Escalation-flagged candidates not selected or consolidated (explicit review required)", BD)
put(wc, r + 5, 3, f'=COUNTIFS($V${R0}:$V${RN},">0",$Y${R0}:$Y${RN},"Consider")+COUNTIFS($V${R0}:$V${RN},">0",$Y${R0}:$Y${RN},"Defer")', fmt="0")

# ================================================================ C_Sensitivity
wsn = wb.create_sheet("C_Sensitivity"); title(wsn, "Section C / Instruction s.20 - Does the selection hold under different weighting schemes?",
                                             "Family score = highest weighted score of any candidate consolidated into that family (so variants do not add up). Ranks recalculate from C_Weights.")
fams = []
for c in rows:
    d, t = DISP[c["id"]]
    key = t if re.fullmatch(r"P\d\d", t) else c["id"]
    if key not in fams: fams.append(key)
fams = sorted(fams, key=lambda k: (0, k) if k.startswith("P") else (1, k))
# helper column in C_Screening: family key
put(wc, 3, 28, "Family key", HDR, fill=HFILL)
for i, c in enumerate(rows):
    r = R0 + i
    put(wc, r, 28, f'=IF(AND(LEFT(Z{r},1)="P",LEN(Z{r})=3),Z{r},A{r})')
wc.column_dimensions["AB"].width = 10
cols = ["Family", "Name / candidate", "Status"] + [f"Score - {k}" for k in WEIGHTS] + [f"Rank - {k}" for k in WEIGHTS] + ["In top 10 under all schemes?", "Selected?", "Consistency"]
header(wsn, 4, cols, [9, 44, 12, 10, 10, 10, 10, 8, 8, 8, 8, 14, 10, 34])
F0 = 5; FN = F0 + len(fams) - 1
fin = {p["id"]: p for p in FINAL}
cname = {c["id"]: c["name"] for c in rows}
for i, k in enumerate(fams):
    r = F0 + i
    put(wsn, r, 1, k); put(wsn, r, 2, fin[k]["name"] if k in fin else cname[k])
    put(wsn, r, 3, "Selected" if k in fin else DISP[k][0])
    for j, sch in enumerate(WEIGHTS):
        col = L(16 + j)
        put(wsn, r, 4 + j, f"=_xlfn.MAXIFS(C_Screening!${col}${R0}:${col}${RN},C_Screening!$AB${R0}:$AB${RN},A{r})", fmt="0.0")
    for j in range(4):
        sc = L(4 + j)
        put(wsn, r, 8 + j, f"=RANK({sc}{r},{sc}${F0}:{sc}${FN})", fmt="0")
    put(wsn, r, 12, f'=IF(MAX(H{r}:K{r})<=10,"Yes",IF(MIN(H{r}:K{r})<=10,"Sometimes","No"))')
    put(wsn, r, 13, "Yes" if k in fin else "No")
    put(wsn, r, 14, f'=IF(AND(M{r}="Yes",L{r}<>"Yes"),"Selected by judgement - see rationale",IF(AND(M{r}="No",L{r}<>"No"),"Scores support selection - explicit review",""))')
wsn.conditional_formatting.add(f"H{F0}:K{FN}", CellIsRule(operator="lessThanOrEqual", formula=["10"], fill=GRF))
wsn.conditional_formatting.add(f"N{F0}:N{FN}", CellIsRule(operator="notEqual", formula=['""'], fill=AMB))
r = FN + 2
notes = [
    "Reading: P01-P04, P06-P09 rank in the top 10 families under every scheme. P05 (integrity) and P10 (AI disclosure) are selected by judgement; C31 (admin sabotage) outranks P10 under all schemes.",
    "P05 is kept as a low-frequency/high-consequence patient-safety scenario (Instruction s.4, s.10). P10 is kept for portfolio coverage and decision usefulness (s.12-13): no other scenario covers negligent insiders or AI, and it drives unique decisions.",
    "C31 is not selected because its treatments (PAM, immutable backup, offboarding) are delivered through P01/P02; management should confirm this reliance (E_Quality, blind spots).",
    "Candidate-level note: a raw top-10 by total score would contain four ransomware-chain variants (C02, C03, C08, C34) and three identity variants (C01, C20, C23) - the variant domination Instruction s.11 and Critical rule 10 warn against.",
]
for i, t in enumerate(notes):
    wsn.cell(row=r + i, column=1, value=t).font = BLK

# ================================================================ D_Portfolio
wd = wb.create_sheet("D_Portfolio"); title(wd, "Section D - Final selected portfolio (10 scenarios)")
dcols = [("id", "ID", 6), ("name", "Scenario name", 30), ("statement", "Scenario statement", 70), ("objective", "Business objective / process", 22), ("assets", "Critical assets", 26),
         ("threat", "Threat source / event", 22), ("vulns", "Vulnerabilities / predisposing conditions", 40), ("consequences", "Principal consequences", 34), ("deps", "Key dependencies", 24),
         ("controls", "Existing controls known", 28), ("appetite", "Risk-appetite / tolerance relevance", 30), ("rationale", "Selection rationale", 50),
         ("evidence", "Evidence basis", 36), ("conf", "Confidence", 10), ("validate", "Information requiring validation", 40), ("cands", "Candidates consolidated", 18), ("prev", "Prior IDs (30-scenario)", 12)]
header(wd, 3, [c[1] for c in dcols], [c[2] for c in dcols])
for i, p in enumerate(FINAL):
    for j, (k, _, _) in enumerate(dcols):
        v = p[k]; v = ", ".join(v) if isinstance(v, list) else v
        put(wd, 4 + i, j + 1, v)
wd.freeze_panes = "C4"

wy = wb.create_sheet("D_WhyTest"); title(wy, "Section D - 'Why this scenario?' test (Instruction s.14)")
header(wy, 3, ["ID", "Scenario"] + WHY_Q, [6, 30, 34, 34, 26, 34, 30, 40])
for i, p in enumerate(FINAL):
    put(wy, 4 + i, 1, p["id"]); put(wy, 4 + i, 2, p["name"])
    for j, a in enumerate(WHY[p["id"]]):
        put(wy, 4 + i, 3 + j, a)

# ---------------------------------------------------------------- D_QuantPrep
wq = wb.create_sheet("D_QuantPrep"); title(wq, "Section D / Instruction s.19 - Preparation for quantitative analysis",
                                          "Preliminary values are taken from the constituent scenarios of the 30-scenario assessment: anchor = lead constituent (core causal chain); range = min-max across constituents. Analyst estimates - validation required.")
S = {s["id"]: s for s in M.SCEN}
PK = [("PbA", "Pb(A)"), ("Pbx", "Pb(psi,A)"), ("De", "de"), ("Dm", "dm"), ("Th", "theta"), ("Mu", "mu(E)")]
cols = ["ID", "Constituents", "Anchor"] + [f"{lab} anchor" for _, lab in PK] + [f"{lab} range" for _, lab in PK] + ["Lowest param. confidence", "CVSS applicability", "Current control effectiveness (FAIR resistance)",
        "FAIR: threat event frequency", "FAIR: threat capability", "FAIR: loss magnitude", "Proposed mitigation (initiatives)", "Prob. reduction (portfolio / 2nd wave)", "Impact reduction (portfolio / 2nd wave)", "Indicative Y1 cost of treatment (CAD)", "Evidence required for validation"]
header(wq, 4, cols, [6, 14, 8] + [8] * 6 + [10] * 6 + [11, 30, 16, 12, 12, 12, 40, 12, 12, 14, 44])
ini = {i[0]: i[1] for i in M.INITIATIVES}
order = {"Low": 0, "Medium": 1, "High": 2}
for i, p in enumerate(FINAL):
    r = 5 + i
    cs = [S[x] for x in p["prev"]]
    anc = cs[0]  # lead constituent = core causal chain
    put(wq, r, 1, p["id"]); put(wq, r, 2, ", ".join(p["prev"])); put(wq, r, 3, anc["id"])
    for j, (k, _) in enumerate(PK):
        put(wq, r, 4 + j, anc["params"][k]["v"], BLUE, "0.00")
        vals = [s["params"][k]["v"] for s in cs]
        put(wq, r, 10 + j, f"{min(vals):.2f}-{max(vals):.2f}" if min(vals) != max(vals) else f"{vals[0]:.2f}")
    confs = [s["params"][k]["conf"] for s in cs for k, _ in PK]
    put(wq, r, 16, min(confs, key=lambda x: order[x]))
    lim = [s["id"] for s in cs if s["cvss_note"].startswith("Limited")]
    put(wq, r, 17, ("Limited for " + ", ".join(lim) + "; " if lim else "Applicable (representative weakness); ") + "anchor CVSS-B " + f"{anc['cvss_score']:.1f}")
    tef, tcap, rs, lm = p["fair"]
    put(wq, r, 18, f"{rs} (weak existing controls)"); put(wq, r, 19, tef); put(wq, r, 20, tcap); put(wq, r, 21, lm)
    inits = sorted({x for s in cs for x in s["inits"]}, key=lambda x: int(x[1:]))
    put(wq, r, 22, "; ".join(f"{x} {ini[x]}" for x in inits))
    enh = BUD.ENHANCED.get(anc["id"], (anc["red_p"], anc["red_i"]))
    put(wq, r, 23, f"{anc['red_p']:.2f} / {enh[0]:.2f}"); put(wq, r, 24, f"{anc['red_i']:.2f} / {enh[1]:.2f}")
    put(wq, r, 25, round(sum(s["cost_y1"] for s in cs), -3), BLUE, "$#,##0")
    put(wq, r, 26, p["validate"])
wq.cell(row=17, column=1, value="Cost = sum of the constituents' allocated Year-1 cost from the 30-scenario model (shared initiatives already split, so no double counting). Re-quantify each P scenario as a whole before management use.").font = BLK

# ---------------------------------------------------------------- D_Validation
wv = wb.create_sheet("D_Validation"); title(wv, "Section D / Instruction s.17 - Cross-functional validation", "R = must validate, C = consult, - = not needed. Anticipated disagreements are recorded, not hidden.")
header(wv, 3, ["ID", "Scenario"] + STAKE + ["Anticipated disagreement"], [6, 30] + [9] * len(STAKE) + [60])
for i, p in enumerate(FINAL):
    r = 4 + i
    put(wv, r, 1, p["id"]); put(wv, r, 2, p["name"])
    for j, v in enumerate(VALID[p["id"]].split()):
        put(wv, r, 3 + j, v, BLUE, fill=(RED if v == "R" else AMB if v == "C" else None))
    put(wv, r, 3 + len(STAKE), DISAGREE[p["id"]])
r = 15
put(wv, r, 2, "Scenarios each stakeholder must validate (R)", BD)
for j in range(len(STAKE)):
    put(wv, r, 3 + j, f'=COUNTIF({L(3+j)}4:{L(3+j)}13,"R")', fmt="0")

# ================================================================ E_Coverage
we = wb.create_sheet("E_Coverage"); title(we, "Section E - Portfolio coverage (Instruction s.12)", "1 = scenario materially represents the category. Categories are included only where a credible material exposure exists.")
header(we, 3, ["Category"] + [p["id"] for p in FINAL] + ["# scenarios", "Status"], [30] + [6] * 10 + [10, 44])
for i, cat in enumerate(CATEGORIES):
    r = 4 + i
    put(we, r, 1, cat)
    for j, p in enumerate(FINAL):
        put(we, r, 2 + j, 1 if cat in COVER[p["id"]] else None, BLUE, "0")
    put(we, r, 12, f"=SUM(B{r}:K{r})", fmt="0")
    note = {"Fraud / BEC": "Not covered - deferred (C18); bounded financial exposure", }.get(cat, "")
    put(we, r, 13, f'=IF(L{r}=0,"Gap - {note or "review"}",IF(L{r}>=8,"Heavily represented - check for variants","Covered"))' if not note else f'=IF(L{r}=0,"Gap - {note}","Covered")')
we.conditional_formatting.add(f"B4:K{3+len(CATEGORIES)}", CellIsRule(operator="equal", formula=["1"], fill=GRF))
we.conditional_formatting.add(f"M4:M{3+len(CATEGORIES)}", FormulaRule(formula=['LEFT(M4,3)="Gap"'], fill=RED))
r = 5 + len(CATEGORIES)
put(we, r, 1, "Categories per scenario", BD)
for j in range(10):
    put(we, r, 2 + j, f"=SUM({L(2+j)}4:{L(2+j)}{3+len(CATEGORIES)})", fmt="0")

# ================================================================ E_Quality
wq2 = wb.create_sheet("E_Quality"); title(wq2, "Section E - Portfolio quality review")
r = table(wq2, 3, ["Question", "Assessment", "Explanation"], [48, 20, 110], QUALITY)
wq2.cell(row=r + 1, column=1, value="Blind spots and watch list").font = SUB
r = table(wq2, r + 2, ["Exposure", "Treatment in this portfolio", "Explanation / trigger to promote"], None, BLIND)
wq2.cell(row=r + 1, column=1, value="Recommended evidence collection before quantification").font = SUB
r = table(wq2, r + 2, ["Evidence activity", "Scenarios", "Owner", "Timing"], None, EVIDENCE_PLAN)
wq2.cell(row=r + 1, column=1, value="Critical-rules check").font = SUB
rules = [("Threat not equated with scenario", "Met - every candidate and final scenario is a causal chain."), ("Not selected on technical severity / CVSS", "Met - CVSS not used in screening."),
         ("Probability not confused with impact", "Met - separate criteria; low-frequency/high-impact retained (P05, P08, P09)."), ("Generic threats not treated as evidence", "Met - evidence basis recorded per candidate; sector-only items marked."),
         ("No double counting of overlapping scenarios", "Met - 24 candidates consolidated into the 10 selected families."), ("No artificial precision", "Met - 1-5 scale; quantitative values flagged as preliminary."),
         ("AI assumptions not presented as facts", "Met - evidence basis and confidence columns."), ("Portfolio not dominated by one technique", "Met - identity and ransomware variants consolidated."),
         ("Business consequences connected", "Met - consequence and objective columns."), ("Selection documented", "Met - rationale and why-test per scenario."),
         ("Assumptions and gaps identified", "Met - A_Context and validation columns."), ("Human accountability", "Pending - requires analyst and cross-functional validation (D_Validation).")]
table(wq2, r + 2, ["Rule", "Status"], None, rules)

# ================================================================ Changes
wch = wb.create_sheet("Changes"); title(wch, "Relationship to the earlier 30-scenario assessment", "The 30-scenario quantification (S1-S30) remains valid as analysis of individual exposures; the P01-P10 portfolio re-groups them for management decision-making.")
header(wch, 3, ["Prior ID", "Prior scenario", "Candidate", "Disposition now", "New family / status"], [8, 60, 10, 14, 50])
cand_of = {}
for c in rows:
    if c["prev"] != "-": cand_of[c["prev"]] = c["id"]
for i, s in enumerate(M.SCEN):
    cid = cand_of[s["id"]]; d, t = DISP[cid]
    put(wch, 4 + i, 1, s["id"]); put(wch, 4 + i, 2, s["name"]); put(wch, 4 + i, 3, cid); put(wch, 4 + i, 4, d); put(wch, 4 + i, 5, t)
r = 5 + len(M.SCEN)
chg = ["Previous 10 (S1-S10) vs new 10: S3 (backup) and S8 (detection) are now stages inside P02 rather than stand-alone scenarios; S1 widened into P01 (identity, incl. S20, S27); S6 widened into P03 (incl. S23); S4 widened into P07 (incl. S11, S17); S9 widened into P08 (incl. S19); S10 widened into P05 (incl. S29).",
       "New in the 10-scenario portfolio: P09 hub single point of failure (S16) and P10 staff disclosure via AI and misdirected communications (S26, S21).",
       "Implication: the CRG KRIs, treatment portfolio and budget were computed on S1-S30 and remain usable; re-quantify P01-P10 as whole scenarios before presenting a 10-scenario register to management."]
for i, t in enumerate(chg):
    wch.cell(row=r + i, column=1, value=t).font = BLK

wb.save(OUT)
print("saved", OUT)
