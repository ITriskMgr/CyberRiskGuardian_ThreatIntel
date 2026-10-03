from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter as L
from openpyxl.formatting.rule import CellIsRule
from data import *

wb = Workbook()
F = "Arial"
BLUE = Font(name=F, color="0000FF", size=10)
BLK = Font(name=F, size=10)
GRN = Font(name=F, color="008000", size=10)
HDR = Font(name=F, bold=True, color="FFFFFF", size=10)
TTL = Font(name=F, bold=True, size=14)
B = Font(name=F, bold=True, size=10)
HFILL = PatternFill("solid", fgColor="1F3864")
YEL = PatternFill("solid", fgColor="FFFF00")
thin = Side(style="thin", color="BFBFBF")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="top")

def header(ws, row, cols, widths=None):
    for j, c in enumerate(cols, 1):
        x = ws.cell(row=row, column=j, value=c)
        x.font, x.fill, x.alignment, x.border = HDR, HFILL, Alignment(wrap_text=True, vertical="center"), BOX
    if widths:
        for j, w in enumerate(widths, 1):
            ws.column_dimensions[L(j)].width = w

def put(ws, r, c, v, font=BLK, fmt=None, wrap=False):
    x = ws.cell(row=r, column=c, value=v)
    x.font, x.border = font, BOX
    if fmt: x.number_format = fmt
    if wrap: x.alignment = WRAP
    return x

# ------------------------------------------------------------------ README
ws = wb.active; ws.title = "README"
ws["A1"] = "MediBec - CyberRiskGuardian Risk Assessment Workbook"; ws["A1"].font = TTL
lines = [
    "Assessment period: next 12 months (October 2026 - September 2027). Currency: CAD. Cost basis: Year-1 = initial implementation + first-year recurring.",
    "Source: MediBec Business Case v2.0b (fictional teaching case, M.-A. Leger) and CyberRiskGuardian Excel Guide v1.0c formulas.",
    "All parameter values, CVSS vectors, costs and reductions are ANALYST ESTIMATES - VALIDATION REQUIRED. Scores are relative decision-support indicators, not loss predictions.",
    "",
    "Colour legend:  blue text = input you may edit  |  black = formula  |  green = link to another sheet  |  yellow fill = key assumption.",
    "",
    "Sheets:",
    "  Parameters  - risk appetite, multiplication factor, tolerance band (edit here).",
    "  Inputs      - S1-S10 quantitative parameters, CVSS v4.0 vector/score, reductions, confidence.",
    "  Analyse     - estimated, tolerated, mitigated, residual risk, ratio, classification, cost-effectiveness (live formulas).",
    "  Portfolio   - 13 initiatives, costs and scenario mapping; shared costs split equally across scenarios addressed (no double-counting).",
    "  Sensitivity - lower / central / higher cases for five key scenarios (all parameters shifted by +/- delta).",
    "  Case_View   - case-native escalation view (s.8 of the case): T x E x I x (1 - C), with the 0.25 / 0.40 escalation thresholds.",
    "  Candidates, Scenarios, KRIs, Evidence - supporting registers.",
    "",
    "Formulas (Excel Guide v1.0c, s.11):",
    "  Estimated = Pb(A) x Pb(psi,A) x CVSS x ((de + dm)/2) x mu(E) / theta x Factor",
    "  Tolerated = Pb(A) x Pb(psi,A) x CVSS x Appetite x mu(E) / theta x Factor",
    "  Mitigated = Estimated x ProbabilityReduction x ImpactReduction ;  Residual = Estimated - Mitigated",
    "  Ratio = Residual / Tolerated ;  Cost-effectiveness = (Estimated - Residual) / Cost x 1000",
    "Note: the template computes Residual as ABS(L - R). Because both reductions are kept within 0-1, L - R is always >= 0 and the two are identical; a validity check column is included.",
    "",
    "How to use: change blue cells on Parameters / Inputs / Portfolio, then recalculate (Formulas > Calculate Now if manual).",
]
for i, t in enumerate(lines, 3):
    ws.cell(row=i, column=1, value=t).font = BLK
ws.column_dimensions["A"].width = 150

# ------------------------------------------------------------------ Parameters
wp = wb.create_sheet("Parameters")
wp["A1"] = "Global parameters"; wp["A1"].font = TTL
rows = [("Risk appetite (analyst estimate)", APPETITE, "Analyst estimate - validation required. Case states 0.40 normalized tolerance; see report s.4."),
        ("Multiplication factor", FACTOR, "Constant across all scenarios (Excel Guide worked example)."),
        ("Tolerance band - lower bound of 'approximately at'", BAND_LOW, "Ratio < this = below tolerance"),
        ("Tolerance band - upper bound of 'approximately at'", BAND_HIGH, "Ratio > this = above tolerance"),
        ("Case appetite (for comparison)", CASE_APPETITE, "Business case s.8"),
        ("Sensitivity delta", 0.10, "Shift applied in Sensitivity sheet"),
        ("Case escalation - executive review threshold", 0.25, "Business case s.8"),
        ("Case escalation - Board acceptance threshold", 0.40, "Business case s.8")]
header(wp, 3, ["Parameter", "Value", "Source / note"], [52, 12, 90])
for i, (a, b, c) in enumerate(rows, 4):
    put(wp, i, 1, a); x = put(wp, i, 2, b, BLUE); put(wp, i, 3, c)
    if i in (4,): x.fill = YEL
APP, FAC, BLO, BHI, CAPP, DELTA, TEXEC, TBOARD = ["Parameters!$B$%d" % r for r in range(4, 12)]

# ------------------------------------------------------------------ Inputs
wi = wb.create_sheet("Inputs")
wi["A1"] = "Scenario inputs (0-1 scale unless stated) - Analyst estimates, validation required"; wi["A1"].font = TTL
cols = ["ID", "Scenario", "Pb(A)", "Pb(psi,A)", "de", "dm", "theta", "mu(E)", "CVSS v4.0 vector", "CVSS-B", "Prob. reduction", "Impact reduction",
        "Conf Pb(A)", "Conf Pb(psi,A)", "Conf de", "Conf dm", "Conf theta", "Conf mu"]
header(wi, 3, cols, [6, 60, 8, 9, 7, 7, 7, 7, 62, 8, 10, 10, 8, 9, 8, 8, 8, 8])
keys = ["PbA", "Pbx", "De", "Dm", "Th", "Mu"]
for r, s in enumerate(SCEN, 4):
    put(wi, r, 1, s["id"]); put(wi, r, 2, s["name"], wrap=True)
    for j, k in enumerate(keys):
        c = put(wi, r, 3 + j, s["params"][k]["v"], BLUE, "0.00")
        c.comment = Comment(s["params"][k]["rat"] + " Evidence: " + s["params"][k]["ev"], "Analyst")
    put(wi, r, 9, s["cvss"], BLUE)
    c = put(wi, r, 10, s["cvss_score"], BLUE, "0.0"); c.comment = Comment(s["cvss_note"][:500], "Analyst")
    put(wi, r, 11, s["red_p"], BLUE, "0.00"); put(wi, r, 12, s["red_i"], BLUE, "0.00")
    for j, k in enumerate(keys):
        put(wi, r, 13 + j, s["params"][k]["conf"], BLUE)
IR0 = 4

# ------------------------------------------------------------------ Portfolio
wpo = wb.create_sheet("Portfolio")
wpo["A1"] = "Recommended treatment portfolio (indicative CAD estimates - not quotations or approved budgets)"; wpo["A1"].font = TTL
pc = ["ID", "Initiative", "Owner", "Priority", "Type", "Initial cost", "Recurring / yr", "Year-1 cost", "# scenarios", "Y1 per scenario"] + [s["id"] for s in SCEN] + ["Start", "End", "Success indicator"]
header(wpo, 3, pc, [6, 44, 26, 11, 16, 12, 12, 12, 9, 12] + [5] * 10 + [9, 9, 50])
P0 = 4
for r, it in enumerate(INITIATIVES, P0):
    iid, name, desc, scs, owner, prio, ini, rec, st, en, succ, dep, typ = it
    put(wpo, r, 1, iid); put(wpo, r, 2, name, wrap=True); put(wpo, r, 3, owner, wrap=True); put(wpo, r, 4, prio); put(wpo, r, 5, typ, wrap=True)
    put(wpo, r, 6, ini, BLUE, "$#,##0"); put(wpo, r, 7, rec, BLUE, "$#,##0")
    put(wpo, r, 8, f"=F{r}+G{r}", fmt="$#,##0")
    put(wpo, r, 9, f"=SUM(K{r}:T{r})", fmt="0")
    put(wpo, r, 10, f"=IF(I{r}=0,0,H{r}/I{r})", fmt="$#,##0")
    for j, s in enumerate(SCEN):
        put(wpo, r, 11 + j, 1 if s["id"] in scs else 0, BLUE, "0")
    put(wpo, r, 21, st); put(wpo, r, 22, en); put(wpo, r, 23, succ, wrap=True)
PN = P0 + len(INITIATIVES) - 1
tr = PN + 1
put(wpo, tr, 2, "TOTAL (each initiative counted once)", B)
for c in (6, 7, 8):
    put(wpo, tr, c, f"=SUM({L(c)}{P0}:{L(c)}{PN})", B, "$#,##0")
for j in range(10):
    c = 11 + j
    put(wpo, tr, c, f"=SUMPRODUCT({L(c)}{P0}:{L(c)}{PN},$J${P0}:$J${PN})", B, "$#,##0")
put(wpo, tr + 1, 2, "Check: sum of scenario allocations - portfolio Year-1 total (should be 0)")
put(wpo, tr + 1, 8, f"=SUM(K{tr}:T{tr})-H{tr}", fmt="$#,##0;($#,##0);-")
wpo.cell(row=tr + 3, column=2, value="Allocation assumption: each initiative's Year-1 cost is split equally across the scenarios it addresses, so shared controls are never charged in full to several scenarios.").font = BLK
wpo.cell(row=tr + 4, column=2, value="Cyber-insurance premium is excluded (Evidence gap - broker quotation required).").font = BLK

# ------------------------------------------------------------------ Analyse
wa = wb.create_sheet("Analyse", 1)
wa["A1"] = "Analyse - CyberRiskGuardian KRIs (Excel Guide v1.0c formulas)"; wa["A1"].font = TTL
wa["A2"] = "Green = linked from Inputs / Portfolio / Parameters; black = formula."; wa["A2"].font = BLK
ac = ["ID", "Scenario", "Pb(A)", "Pb(psi,A)", "CVSS-B", "de", "dm", "mu(E)", "theta", "Appetite", "Factor",
      "Estimated risk", "Tolerated risk", "Pre-treatment ratio", "Allocated Y1 cost", "Prob. reduction", "Impact reduction",
      "Mitigated", "Residual", "Residual / Tolerated", "Classification", "Cost-effectiveness (per $1k)", "Validity (P,Q in 0-1)"]
header(wa, 4, ac, [6, 52, 7, 8, 7, 6, 6, 7, 7, 8, 8, 11, 11, 10, 12, 9, 9, 11, 11, 10, 24, 12, 10])
A0 = 5
for idx, s in enumerate(SCEN):
    r = A0 + idx; ir = IR0 + idx
    put(wa, r, 1, f"=Inputs!A{ir}", GRN); put(wa, r, 2, f"=Inputs!B{ir}", GRN, wrap=True)
    put(wa, r, 3, f"=Inputs!C{ir}", GRN, "0.00"); put(wa, r, 4, f"=Inputs!D{ir}", GRN, "0.00")
    put(wa, r, 5, f"=Inputs!J{ir}", GRN, "0.0"); put(wa, r, 6, f"=Inputs!E{ir}", GRN, "0.00")
    put(wa, r, 7, f"=Inputs!F{ir}", GRN, "0.00"); put(wa, r, 8, f"=Inputs!H{ir}", GRN, "0.00")
    put(wa, r, 9, f"=Inputs!G{ir}", GRN, "0.00"); put(wa, r, 10, f"={APP}", GRN, "0.00"); put(wa, r, 11, f"={FAC}", GRN, "#,##0")
    put(wa, r, 12, f"=C{r}*D{r}*E{r}*((F{r}+G{r})/2)*H{r}/I{r}*K{r}", fmt="#,##0")
    put(wa, r, 13, f"=C{r}*D{r}*E{r}*J{r}*H{r}/I{r}*K{r}", fmt="#,##0")
    put(wa, r, 14, f"=L{r}/M{r}", fmt="0.00")
    col = L(11 + idx)
    put(wa, r, 15, f"=Portfolio!{col}{tr}", GRN, "$#,##0")
    put(wa, r, 16, f"=Inputs!K{ir}", GRN, "0.00"); put(wa, r, 17, f"=Inputs!L{ir}", GRN, "0.00")
    put(wa, r, 18, f"=L{r}*P{r}*Q{r}", fmt="#,##0")
    put(wa, r, 19, f"=L{r}-R{r}", fmt="#,##0")
    put(wa, r, 20, f"=S{r}/M{r}", fmt="0.00")
    put(wa, r, 21, f'=IF(T{r}<{BLO},"Below tolerance",IF(T{r}<={BHI},"Approximately at tolerance","Above tolerance"))')
    put(wa, r, 22, f"=IF(O{r}=0,0,(L{r}-S{r})/O{r}*1000)", fmt="0.00")
    put(wa, r, 23, f'=IF(AND(P{r}>=0,P{r}<=1,Q{r}>=0,Q{r}<=1),"OK","CHECK")')
AN = A0 + len(SCEN) - 1
t = AN + 1
put(wa, t, 2, "Portfolio totals", B)
for c in (12, 13, 18, 19):
    put(wa, t, c, f"=SUM({L(c)}{A0}:{L(c)}{AN})", B, "#,##0")
put(wa, t, 15, f"=SUM(O{A0}:O{AN})", B, "$#,##0")
put(wa, t, 20, f"=S{t}/M{t}", B, "0.00")
put(wa, t, 22, f"=(L{t}-S{t})/O{t}*1000", B, "0.00")
red = PatternFill("solid", fgColor="F4CCCC"); amb = PatternFill("solid", fgColor="FFF2CC"); gr = PatternFill("solid", fgColor="D9EAD3")
rng = f"T{A0}:T{AN}"
wa.conditional_formatting.add(rng, CellIsRule(operator="greaterThan", formula=[BHI], fill=red))
wa.conditional_formatting.add(rng, CellIsRule(operator="between", formula=[BLO, BHI], fill=amb))
wa.conditional_formatting.add(rng, CellIsRule(operator="lessThan", formula=[BLO], fill=gr))
wa.freeze_panes = "C5"
wa.cell(row=t + 2, column=2, value="Interpretation: Residual/Tolerated simplifies to ((de+dm)/2) x (1 - P x Q) / Appetite - CVSS, probabilities, resilience and criticality scale magnitude (ranking) but cancel out of the tolerance test.").font = BLK

# ------------------------------------------------------------------ Sensitivity
wsn = wb.create_sheet("Sensitivity")
wsn["A1"] = "Sensitivity - lower / central / higher cases"; wsn["A1"].font = TTL
wsn["A2"] = "Lower case: Pb(A), Pb(psi,A), de, dm -delta; theta +delta; reductions +delta. Higher case: the reverse. Values clamped to [0.05, 0.99]."; wsn["A2"].font = BLK
sc = ["ID", "Case", "Shift", "Pb(A)", "Pb(psi,A)", "de", "dm", "theta", "mu", "CVSS", "P red", "I red", "Estimated", "Tolerated", "Residual", "Ratio", "Classification"]
header(wsn, 4, sc, [6, 10, 7, 7, 8, 7, 7, 7, 7, 7, 7, 7, 11, 11, 11, 8, 26])
r = 5
for sid in SENS_IDS:
    idx = [s["id"] for s in SCEN].index(sid); ir = IR0 + idx
    for lab, sign in (("Lower", -1), ("Central", 0), ("Higher", 1)):
        put(wsn, r, 1, sid); put(wsn, r, 2, lab)
        put(wsn, r, 3, f"={sign}*{DELTA}", fmt="0.00")
        cl = lambda e: f"=MAX(0.05,MIN(0.99,{e}))"
        put(wsn, r, 4, cl(f"Inputs!C{ir}+C{r}"), fmt="0.00"); put(wsn, r, 5, cl(f"Inputs!D{ir}+C{r}"), fmt="0.00")
        put(wsn, r, 6, cl(f"Inputs!E{ir}+C{r}"), fmt="0.00"); put(wsn, r, 7, cl(f"Inputs!F{ir}+C{r}"), fmt="0.00")
        put(wsn, r, 8, cl(f"Inputs!G{ir}-C{r}"), fmt="0.00"); put(wsn, r, 9, f"=Inputs!H{ir}", GRN, "0.00")
        put(wsn, r, 10, f"=Inputs!J{ir}", GRN, "0.0")
        put(wsn, r, 11, cl(f"Inputs!K{ir}-C{r}"), fmt="0.00"); put(wsn, r, 12, cl(f"Inputs!L{ir}-C{r}"), fmt="0.00")
        put(wsn, r, 13, f"=D{r}*E{r}*J{r}*((F{r}+G{r})/2)*I{r}/H{r}*{FAC}", fmt="#,##0")
        put(wsn, r, 14, f"=D{r}*E{r}*J{r}*{APP}*I{r}/H{r}*{FAC}", fmt="#,##0")
        put(wsn, r, 15, f"=M{r}-M{r}*K{r}*L{r}", fmt="#,##0")
        put(wsn, r, 16, f"=O{r}/N{r}", fmt="0.00")
        put(wsn, r, 17, f'=IF(P{r}<{BLO},"Below tolerance",IF(P{r}<={BHI},"Approximately at tolerance","Above tolerance"))')
        r += 1
wsn.conditional_formatting.add(f"P5:P{r-1}", CellIsRule(operator="greaterThan", formula=[BHI], fill=red))
wsn.conditional_formatting.add(f"P5:P{r-1}", CellIsRule(operator="between", formula=[BLO, BHI], fill=amb))
wsn.conditional_formatting.add(f"P5:P{r-1}", CellIsRule(operator="lessThan", formula=[BLO], fill=gr))
r += 1
wsn.cell(row=r, column=1, value="Appetite sensitivity (central parameters): ratio at alternative appetite values").font = B
header(wsn, r + 1, ["ID", "Appetite 0.20", "Appetite 0.30", "Appetite 0.40"])
for i, s in enumerate(SCEN):
    rr = r + 2 + i; ar = A0 + i
    put(wsn, rr, 1, s["id"])
    for j, a in enumerate((0.2, 0.3, 0.4)):
        put(wsn, rr, 2 + j, f"=Analyse!S{ar}/(Analyse!M{ar}/Analyse!J{ar}*{a})", fmt="0.00")

# ------------------------------------------------------------------ Case view
wc = wb.create_sheet("Case_View")
wc["A1"] = "Case-native escalation view (Business case s.8) - secondary cross-check"; wc["A1"].font = TTL
wc["A2"] = "Normalized residual = T x E x I x (1 - C), with T = Pb(A), E = Pb(psi,A), I = (de+dm)/2, C (control maturity) proxied by theta. Post-treatment = current x (1 - P x Q). Analyst mapping - validate."; wc["A2"].font = BLK
header(wc, 4, ["ID", "T", "E", "I", "C (theta)", "Current residual", "Level", "Post-treatment residual", "Level", "Acceptance authority (post)"], [6, 7, 7, 7, 9, 12, 11, 14, 11, 28])
lv = lambda c: f'=IF({c}<=0.07,"Low",IF({c}<=0.15,"Moderate",IF({c}<=0.25,"High","Critical")))'
for i, s in enumerate(SCEN):
    r = 5 + i; ir = IR0 + i
    put(wc, r, 1, s["id"])
    put(wc, r, 2, f"=Inputs!C{ir}", GRN, "0.00"); put(wc, r, 3, f"=Inputs!D{ir}", GRN, "0.00")
    put(wc, r, 4, f"=(Inputs!E{ir}+Inputs!F{ir})/2", fmt="0.00"); put(wc, r, 5, f"=Inputs!G{ir}", GRN, "0.00")
    put(wc, r, 6, f"=B{r}*C{r}*D{r}*(1-E{r})", fmt="0.000"); put(wc, r, 7, lv(f"F{r}"))
    put(wc, r, 8, f"=F{r}*(1-Inputs!K{ir}*Inputs!L{ir})", fmt="0.000"); put(wc, r, 9, lv(f"H{r}"))
    put(wc, r, 10, f'=IF(H{r}>{TBOARD},"Board / ownership group",IF(H{r}>{TEXEC},"Executive leadership","CIO / Cybersecurity lead"))')

# ------------------------------------------------------------------ registers
def register(title, cols, widths, rows):
    w = wb.create_sheet(title)
    header(w, 1, cols, widths)
    for i, row in enumerate(rows, 2):
        for j, v in enumerate(row, 1):
            put(w, i, j, v, wrap=True)
    w.freeze_panes = "A2"
    return w

register("Scenarios", ["ID", "Library ref", "Scenario", "Statement", "Threat source", "Vulnerabilities / conditions", "Existing controls (evidence)", "Treatment package", "Initiatives", "Risk owner", "Horizon", "CVSS rationale"],
         [6, 14, 34, 70, 24, 50, 40, 50, 20, 22, 12, 60],
         [(s["id"], s["ref"], s["name"], s["statement"], s["threat_source"], "; ".join(s["vulns"]), "; ".join(s["controls"]), s["treatment"], ", ".join(s["inits"]), s["owner"], s["horizon"], s["cvss_note"]) for s in SCEN])
register("Candidates", ["ID", "Candidate scenario", "Threat source", "Vulnerability / condition", "Primary asset", "Undesired outcome", "Library ref", "Decision", "Screening rationale"],
         [6, 44, 22, 38, 14, 28, 14, 18, 50], CANDIDATES)
register("KRIs", ["ID", "KRI", "Scenarios", "Measurement", "Data source", "Owner", "Frequency", "Target", "Warning", "Critical"],
         [8, 44, 12, 34, 24, 20, 11, 12, 12, 14], KRIS)
register("Evidence", ["ID", "Source", "Basis", "Content used"], [6, 60, 30, 100], EVIDENCE)

for w in wb.worksheets:
    w.sheet_view.showGridLines = False
wb.save("../output/MediBec_CyberRiskGuardian_Workbook.xlsx")
print("saved")
