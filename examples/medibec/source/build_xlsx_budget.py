import json
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.formatting.rule import CellIsRule
import data30 as M

B = json.load(open("budget.json"))
SRC = "../output/MediBec_CyberRiskGuardian_Workbook_30.xlsx"
OUT = "../output/MediBec_CyberRiskGuardian_Workbook_30_Budget.xlsx"
wb = load_workbook(SRC)

F = "Arial"
BLUE = Font(name=F, color="0000FF", size=10); BLK = Font(name=F, size=10); GRN = Font(name=F, color="008000", size=10)
HDR = Font(name=F, bold=True, color="FFFFFF", size=10); TTL = Font(name=F, bold=True, size=14); BD = Font(name=F, bold=True, size=10)
HFILL = PatternFill("solid", fgColor="1F3864"); YEL = PatternFill("solid", fgColor="FFFF00")
thin = Side(style="thin", color="BFBFBF"); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="top")
red = PatternFill("solid", fgColor="F4CCCC"); amb = PatternFill("solid", fgColor="FFF2CC"); gr = PatternFill("solid", fgColor="D9EAD3")
MONEY = "$#,##0;($#,##0);-"

def header(ws, row, cols, widths=None):
    for j, c in enumerate(cols, 1):
        x = ws.cell(row=row, column=j, value=c); x.font, x.fill, x.border = HDR, HFILL, BOX
        x.alignment = Alignment(wrap_text=True, vertical="center")
    if widths:
        for j, w in enumerate(widths, 1): ws.column_dimensions[L(j)].width = w
def put(ws, r, c, v, font=BLK, fmt=None, wrap=False):
    x = ws.cell(row=r, column=c, value=v); x.font, x.border = font, BOX
    if fmt: x.number_format = fmt
    if wrap: x.alignment = WRAP
    return x

# ------------------------------------------------------------------ Budget_Params
wp = wb.create_sheet("Budget_Params")
wp["A1"] = "Cybersecurity budget guideline and appetite mapping"; wp["A1"].font = TTL
header(wp, 3, ["Parameter", "Value", "Source / note"], [58, 16, 100])
params = [
    ("Total IT budget incl. salaries (CAD)", B["IT_BUDGET"], MONEY, "Business case s.4 (E3). Evidence gap - appears high for a 15-person IT team; validate with Finance."),
    ("Guideline floor - risk seeking (% of IT budget)", B["BAND"]["min"], "0.0%", "User-supplied guideline."),
    ("Guideline median - risk neutral", B["BAND"]["median"], "0.0%", "User-supplied guideline."),
    ("Guideline ceiling - risk averse", B["BAND"]["max"], "0.0%", "User-supplied guideline."),
    ("Risk appetite (linked)", "=Parameters!B4", "0.00", "From Parameters sheet (analyst estimate 0.30)."),
    ("Appetite anchor for ceiling (risk averse)", 0.30, "0.00", "Appetite scale anchor 0.3 = risk averse."),
    ("Appetite anchor for median (risk neutral)", 0.50, "0.00", "Appetite scale anchor 0.5 = risk neutral."),
    ("Appetite anchor for floor (risk seeking)", 0.70, "0.00", "Appetite scale anchor 0.7 = relatively high appetite."),
    ("Contingency and incident reserve (% of cyber spend)", B["RESERVE_PCT"], "0%", "Analyst convention."),
    ("Recommended ramp - Year 1 (% of IT budget)", B["RAMP"][0], "0.0%", "Delivery capacity limits Year 1."),
    ("Recommended ramp - Year 2", B["RAMP"][1], "0.0%", ""),
    ("Recommended ramp - Year 3", B["RAMP"][2], "0.0%", "Reaches the appetite-consistent target."),
    ("FTE annual value (CAD)", B["FTE"], MONEY, "Workbook convention 1,950 h x $80."),
]
for i, (a, v, f, n) in enumerate(params, 4):
    put(wp, i, 1, a); x = put(wp, i, 2, v, GRN if isinstance(v, str) else BLUE, f); put(wp, i, 3, n, wrap=True)
    if i == 4: x.fill = YEL
ITB, BMIN, BMED, BMAX, APP, A1, A2, A3, CONT, R1, R2, R3, FTEC = ["Budget_Params!$B$%d" % r for r in range(4, 17)]
r = 18
put(wp, r, 1, "Appetite-consistent target (% of IT budget)", BD)
tgt = (f"=IF({APP}<={A1},{BMAX},IF({APP}<={A2},{BMAX}+({APP}-{A1})/({A2}-{A1})*({BMED}-{BMAX}),"
       f"IF({APP}<={A3},{BMED}+({APP}-{A2})/({A3}-{A2})*({BMIN}-{BMED}),{BMIN})))")
put(wp, r, 2, tgt, BD, "0.0%"); put(wp, r, 3, "Piecewise-linear: 0.30 -> 12%, 0.50 -> 7.8%, 0.70 -> 4%; clamped outside.")
put(wp, r + 1, 1, "Appetite-consistent target (CAD per year)", BD); put(wp, r + 1, 2, f"=B{r}*{ITB}", BD, MONEY)
TGT, TGTC = f"Budget_Params!$B${r}", f"Budget_Params!$B${r+1}"
put(wp, r + 3, 1, "Target by appetite value", BD)
header(wp, r + 4, ["Appetite", "Target %", "Target CAD"])
for i, a in enumerate((0.2, 0.3, 0.4, 0.5, 0.6, 0.7, 0.8)):
    rr = r + 5 + i
    put(wp, rr, 1, a, BLUE, "0.00")
    put(wp, rr, 2, tgt.replace(APP, f"A{rr}"), fmt="0.0%")
    put(wp, rr, 3, f"=B{rr}*{ITB}", fmt=MONEY)
wp.column_dimensions["A"].width = 58

# ------------------------------------------------------------------ Budget_Lines
wl = wb.create_sheet("Budget_Lines")
wl["A1"] = "Cybersecurity and IT budget lines - fiscal years Oct-Sep (Y1 = Oct 2026-Sep 2027)"; wl["A1"].font = TTL
wl["A2"] = "Gross = total cost of the line (phased: initial spread over start-end months, recurring from completion). Cyber share splits it between the cybersecurity budget and the IT project/operating budget. Blue = editable input."; wl["A2"].font = BLK
cols = ["Group", "ID", "Line", "Type", "Cyber share", "Gross Y1", "Gross Y2", "Gross Y3", "Cyber Y1", "Cyber Y2", "Cyber Y3", "IT Y1", "IT Y2", "IT Y3", "Note"]
header(wl, 4, cols, [34, 7, 46, 9, 9, 12, 12, 12, 12, 12, 12, 12, 12, 12, 60])
row = 5
share_of = dict(B["CYBER_SHARE"]); share_of.update({e[0]: e[7] for e in B["SECOND"]})
def line(group, iid, name, typ, share, gross, note=""):
    global row
    put(wl, row, 1, group); put(wl, row, 2, iid); put(wl, row, 3, name, wrap=True); put(wl, row, 4, typ)
    put(wl, row, 5, share, BLUE, "0%")
    for j in range(3):
        put(wl, row, 6 + j, gross[j], BLUE, MONEY)
        put(wl, row, 9 + j, f"={L(6+j)}{row}*$E{row}", fmt=MONEY)
        put(wl, row, 12 + j, f"={L(6+j)}{row}*(1-$E{row})", fmt=MONEY)
    put(wl, row, 15, note, wrap=True)
    row += 1
for l in B["lines"]:
    g = [a + b for a, b in zip(l["cyber"], l["it"])]
    if sum(g) == 0: continue
    sh = 1.0 if l["group"][0] in "AD" else share_of[l["id"]]
    note = B["SHARE_WHY"].get(l["id"], "") if l["group"][0] in "BC" else ""
    line(l["group"], l["id"], l["name"], l["kind"], sh, g, note)
T2G = "G. Tranche 2 - appetite alignment"
for h in B["T2_ROLES"]:
    import budget as BB
    y = BB.years(BB.phase(0, B["FTE"], h[2], h[2]))
    line(T2G, h[0], h[1], "run", 1.0, y, f"Hire month {h[2]}")
for t in B["T2_FIXED"]:
    line(T2G, t[0], t[1], "run" if t[0] == "T2-2" else "project", 1.0, t[3], t[4])
LAST = row - 1
# contingency (10% of groups A-F cyber) and reserve (balancing to ramp target)
put(wl, row, 1, "H. Contingency and incident reserve"); put(wl, row, 2, "R1"); put(wl, row, 3, "Contingency (overruns, incident response surge)"); put(wl, row, 4, "run"); put(wl, row, 5, 1.0, BLUE, "0%")
for j in range(3):
    c = L(9 + j)
    put(wl, row, 9 + j, f'={CONT}*SUMPRODUCT((LEFT($A$5:$A${LAST},1)<="F")*{c}5:{c}{LAST})', fmt=MONEY)
    put(wl, row, 6 + j, f"={c}{row}", fmt=MONEY); put(wl, row, 12 + j, 0, fmt=MONEY)
CONT_ROW = row; row += 1
put(wl, row, 1, "I. Risk-reduction investment reserve"); put(wl, row, 2, "R2"); put(wl, row, 3, "Reserve released by Steering Committee against validated, scenario-linked business cases"); put(wl, row, 4, "project"); put(wl, row, 5, 1.0, BLUE, "0%")
for j, rr in enumerate((R1, R2, R3)):
    c = L(9 + j)
    put(wl, row, 9 + j, f"=MAX(0,{rr}*{ITB}-SUM({c}5:{c}{CONT_ROW}))", fmt=MONEY)
    put(wl, row, 6 + j, f"={c}{row}", fmt=MONEY); put(wl, row, 12 + j, 0, fmt=MONEY)
put(wl, row, 15, "Balancing line: brings the recommended budget to the ramp target. Not spent unless a business case tied to scenario IDs is approved.", wrap=True)
RES_ROW = row; row += 1
put(wl, row, 3, "TOTAL", BD)
for j in range(9):
    c = L(6 + j); put(wl, row, 6 + j, f"=SUM({c}5:{c}{RES_ROW})", BD, MONEY)
TOT_ROW = row
wl.freeze_panes = "D5"

# ------------------------------------------------------------------ Budget_Summary
ws = wb.create_sheet("Budget_Summary")
ws["A1"] = "Cybersecurity budget options vs. guideline band (4% - 12% of IT budget)"; ws["A1"].font = TTL
header(ws, 3, ["Option", "Groups included", "Cyber Y1", "Cyber Y2", "Cyber Y3", "% IT Y1", "% IT Y2", "% IT Y3", "Band check (all years)", "Position vs appetite target (Y3)", "Implied appetite (Y3)"],
       [44, 18, 13, 13, 13, 9, 9, 9, 22, 28, 12])
def sumg(letters, yr):
    c = L(9 + yr)
    return "+".join(f'SUMPRODUCT((LEFT(Budget_Lines!$A$5:$A${RES_ROW},1)="{g}")*Budget_Lines!{c}5:{c}{RES_ROW})' for g in letters)
opts = [("A. Status quo (baseline only)", "A", False), ("B. Baseline + 30-scenario portfolio", "ABC", False),
        ("C. B + security team expansion (+10% contingency)", "ABCD", True), ("D. C + second-wave treatment (+10% contingency)", "ABCDEF", True),
        ("R. Recommended: D + Tranche 2 + risk-reduction reserve", "ABCDEFGHI", False)]
for i, (lab, gs, cont) in enumerate(opts):
    r = 4 + i
    put(ws, r, 1, lab, BD if lab.startswith("R.") else BLK); put(ws, r, 2, gs)
    for y in range(3):
        f = sumg(gs, y)
        put(ws, r, 3 + y, f"=({f})*(1+{CONT})" if cont else f"={f}", fmt=MONEY)
        put(ws, r, 6 + y, f"={L(3+y)}{r}/{ITB}", fmt="0.0%")
    put(ws, r, 9, f'=IF(MIN(F{r}:H{r})<{BMIN},"Below floor",IF(MAX(F{r}:H{r})>{BMAX},"Above ceiling","Within band"))')
    put(ws, r, 10, f'=IF(ABS(H{r}-{TGT})<=0.005,"At appetite target",IF(H{r}<{TGT},"Below target by "&TEXT({TGT}-H{r},"0.0%"),"Above target by "&TEXT(H{r}-{TGT},"0.0%")))')
    put(ws, r, 11, f'=IF(H{r}<{BMIN},"n/a (below floor)",IF(H{r}>={BMAX},{A1},IF(H{r}>={BMED},{A1}+({BMAX}-H{r})/({BMAX}-{BMED})*({A2}-{A1}),{A2}+({BMED}-H{r})/({BMED}-{BMIN})*({A3}-{A2}))))', fmt="0.00")
ws.conditional_formatting.add("I4:I8", CellIsRule(operator="equal", formula=['"Within band"'], fill=gr))
ws.conditional_formatting.add("I4:I8", CellIsRule(operator="notEqual", formula=['"Within band"'], fill=red))
r = 11
put(ws, r, 1, "Appetite-consistent target", BD); put(ws, r, 3, f"={TGTC}", BD, MONEY); put(ws, r, 6, f"={TGT}", BD, "0.0%")
put(ws, r + 1, 1, "Check: recommended Y3 equals target (0 = OK)"); put(ws, r + 1, 5, f"=ROUND(E8-{TGTC},0)", fmt=MONEY)

r = 15
ws.cell(row=r, column=1, value="Project budgets - recommended option (one-time spend)").font = BD
header(ws, r + 1, ["Budget", "", "Y1", "Y2", "Y3", "3-year total"])
proj = [("Cybersecurity projects (cyber share)", "cyber"), ("Security-relevant IT projects (IT share)", "it")]
for i, (lab, kind) in enumerate(proj):
    rr = r + 2 + i; put(ws, rr, 1, lab)
    for y in range(3):
        c = L((9 if kind == "cyber" else 12) + y)
        put(ws, rr, 3 + y, f'=SUMPRODUCT((Budget_Lines!$D$5:$D${RES_ROW}="project")*Budget_Lines!{c}5:{c}{RES_ROW})', fmt=MONEY)
    put(ws, rr, 6, f"=SUM(C{rr}:E{rr})", fmt=MONEY)
rr = r + 4; put(ws, rr, 1, "Total IT and cybersecurity projects", BD)
for y in range(4):
    put(ws, rr, 3 + y, f"={L(3+y)}{r+2}+{L(3+y)}{r+3}", BD, MONEY)
rr = r + 5; put(ws, rr, 1, "Cybersecurity run (operating) budget")
for y in range(3):
    put(ws, rr, 3 + y, f"=C8-C{r+2}".replace("C8", f"{L(3+y)}8").replace(f"C{r+2}", f"{L(3+y)}{r+2}"), fmt=MONEY)
put(ws, rr, 6, f"=SUM(C{rr}:E{rr})", fmt=MONEY)

r = 23
ws.cell(row=r, column=1, value="Sensitivity to the IT budget figure (evidence gap) - option D and recommended").font = BD
header(ws, r + 1, ["IT budget (CAD)", "", "D % Y1", "D % Y2", "D % Y3", "D band check", "Target CAD @ appetite"])
for i, itb in enumerate((100e6, 75e6, 50e6, 25e6, 10e6)):
    rr = r + 2 + i
    put(ws, rr, 1, itb, BLUE, MONEY)
    for y in range(3):
        put(ws, rr, 3 + y, f"={L(3+y)}7/A{rr}", fmt="0.0%")
    put(ws, rr, 6, f'=IF(MIN(C{rr}:E{rr})<{BMIN},"Below floor",IF(MAX(C{rr}:E{rr})>{BMAX},"Above ceiling - phase or cut","Within band"))')
    put(ws, rr, 7, f"={TGT}*A{rr}", fmt=MONEY)

# ------------------------------------------------------------------ Budget_Risk
wr = wb.create_sheet("Budget_Risk")
wr["A1"] = "Residual risk by funding level (same CRG formulas; second-wave reductions are analyst estimates)"; wr["A1"].font = TTL
header(wr, 3, ["ID", "Scenario", "Estimated", "Tolerated", "P (portfolio)", "Q (portfolio)", "Residual (portfolio)", "Ratio (portfolio)", "Status (portfolio)",
               "P (with 2nd wave)", "Q (with 2nd wave)", "Residual (2nd wave)", "Ratio (2nd wave)", "Status (2nd wave)"],
       [6, 50, 10, 10, 9, 9, 11, 9, 22, 9, 9, 11, 9, 22])
cls = lambda c: f'=IF({c}<Parameters!$B$6,"Below tolerance",IF({c}<=Parameters!$B$7,"Approximately at tolerance","Above tolerance"))'
for i, s in enumerate(M.SCEN):
    r = 4 + i; ar = 5 + i; ir = 4 + i
    put(wr, r, 1, f"=Analyse!A{ar}", GRN); put(wr, r, 2, f"=Analyse!B{ar}", GRN, wrap=True)
    put(wr, r, 3, f"=Analyse!L{ar}", GRN, "#,##0"); put(wr, r, 4, f"=Analyse!M{ar}", GRN, "#,##0")
    put(wr, r, 5, f"=Inputs!K{ir}", GRN, "0.00"); put(wr, r, 6, f"=Inputs!L{ir}", GRN, "0.00")
    put(wr, r, 7, f"=C{r}*(1-E{r}*F{r})", fmt="#,##0"); put(wr, r, 8, f"=G{r}/D{r}", fmt="0.00"); put(wr, r, 9, cls(f"H{r}"))
    ep, eq = B["ENHANCED"].get(s["id"], (s["red_p"], s["red_i"]))
    put(wr, r, 10, ep, BLUE, "0.00"); put(wr, r, 11, eq, BLUE, "0.00")
    put(wr, r, 12, f"=C{r}*(1-J{r}*K{r})", fmt="#,##0"); put(wr, r, 13, f"=L{r}/D{r}", fmt="0.00"); put(wr, r, 14, cls(f"M{r}"))
last = 3 + len(M.SCEN)
for col in ("H", "M"):
    rng = f"{col}4:{col}{last}"
    wr.conditional_formatting.add(rng, CellIsRule(operator="greaterThan", formula=["Parameters!$B$7"], fill=red))
    wr.conditional_formatting.add(rng, CellIsRule(operator="between", formula=["Parameters!$B$6", "Parameters!$B$7"], fill=amb))
    wr.conditional_formatting.add(rng, CellIsRule(operator="lessThan", formula=["Parameters!$B$6"], fill=gr))
r = last + 2
put(wr, r, 2, "Count - above / approximately at / below", BD)
for j, (cc, lab) in enumerate((("I", "portfolio"), ("N", "2nd wave"))):
    put(wr, r + 1 + j, 2, f"Funding level: {lab}")
    for k, t in enumerate(("Above tolerance", "Approximately at tolerance", "Below tolerance")):
        put(wr, r + 1 + j, 3 + k, f'=COUNTIF({cc}4:{cc}{last},"{t}")', fmt="0")
put(wr, r + 3, 2, "Total residual (portfolio / 2nd wave)")
put(wr, r + 3, 3, f"=SUM(G4:G{last})", fmt="#,##0"); put(wr, r + 3, 4, f"=SUM(L4:L{last})", fmt="#,##0")

# ------------------------------------------------------------------ KRIs + README
wk = wb["KRIs"]
kr = wk.max_row + 1
for i, k in enumerate(B["BUDGET_KRIS"]):
    for j, v in enumerate(k, 1):
        put(wk, kr + i, j, v, wrap=True)
rd = wb["README"]
n = rd.max_row + 2
for t in ["Budget extension (added):",
          "  Budget_Params  - IT budget, guideline band (4% / 7.8% / 12%), appetite-to-target mapping, ramp.",
          "  Budget_Lines   - phased cyber and IT lines per fiscal year with cyber share; contingency and balancing reserve.",
          "  Budget_Summary - options A-D and recommended vs band; implied appetite; project budgets; IT-budget sensitivity.",
          "  Budget_Risk    - residual risk and tolerance status at portfolio vs. second-wave funding."]:
    rd.cell(row=n, column=1, value=t).font = BLK; n += 1
for w in (wp, wl, ws, wr):
    w.sheet_view.showGridLines = False
wb.save(OUT)
print("saved", RES_ROW, TOT_ROW)
